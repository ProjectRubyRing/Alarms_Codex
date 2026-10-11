"""Render workbook cell settings for review. This is not an Excel screenshot."""
from pathlib import Path
import math
from PIL import Image, ImageDraw, ImageFont
import openpyxl
from build_guide import XLSX, VALIDATION, FONT_PATH, FONT_BOLD, wrapped

def color(value, default='#FFFFFF'):
    if value is not None and value.type=='rgb' and isinstance(value.rgb,str):
        return '#'+value.rgb[-6:]
    return default

def dimensions(ws):
    widths=[]
    for i in range(1,ws.max_column+1):
        letter=openpyxl.utils.get_column_letter(i)
        width=ws.column_dimensions[letter].width
        widths.append(int(width*7+.5))
    heights=[int((ws.row_dimensions[i].height or 15)*96/72+.5) for i in range(1,ws.max_row+1)]
    return widths,heights

def merges(ws):
    origins={}; covered=set()
    for m in ws.merged_cells.ranges:
        origins[(m.min_row,m.min_col)]=m
        for row in range(m.min_row,m.max_row+1):
            for col in range(m.min_col,m.max_col+1):
                if (row,col)!=(m.min_row,m.min_col): covered.add((row,col))
    return origins,covered

def font(cell):
    size=max(1,round(float(cell.font.sz or 10.5)*96/72))
    return ImageFont.truetype(FONT_BOLD if cell.font.bold else FONT_PATH,size,index=2)

def cell_geometry(ws,r,c,widths,heights,origins):
    merged=origins.get((r,c))
    rr=merged.max_row if merged else r
    cc=merged.max_col if merged else c
    return sum(widths[c-1:cc]),sum(heights[r-1:rr])

def line_metrics(cell,text,width):
    f=font(cell)
    lines=wrapped(text,f,max(5,width-12)) if cell.alignment.wrap_text else str(text).split('\n')
    asc,desc=f.getmetrics()
    return f,lines,asc+desc

def inspect_layout(workbook):
    issues=[]
    for ws in workbook:
        widths,heights=dimensions(ws)
        origins,covered=merges(ws)
        for rr in ws:
            for cell in rr:
                if cell.value is None or (cell.row,cell.column) in covered: continue
                width,height=cell_geometry(ws,cell.row,cell.column,widths,heights,origins)
                f,lines,lh=line_metrics(cell,str(cell.value),width)
                if len(lines)*lh+4>height+2:
                    issues.append({'sheet':ws.title,'cell':cell.coordinate,'height':height,'needed':len(lines)*lh+4,
                                   'text':str(cell.value)[:65]})
                if not cell.alignment.wrap_text and max(f.getlength(x) for x in lines)>width-6:
                    issues.append({'sheet':ws.title,'cell':cell.coordinate,'width':width,'text':str(cell.value)[:65]})
    return issues

def render(ws,name,start=1,end=20):
    widths,heights=dimensions(ws)
    origins,covered=merges(ws)
    y0=sum(heights[:start-1])
    canvas=Image.new('RGB',(sum(widths)+24,sum(heights[start-1:end])+58),'#FFFFFF')
    draw=ImageDraw.Draw(canvas)
    label_font=ImageFont.truetype(FONT_PATH,12,index=2)
    draw.text((12,8),'セル設定に基づく描画プレビュー（Excel実機画面ではありません）',font=label_font,fill='#555555')
    for r in range(start,end+1):
        for c in range(1,ws.max_column+1):
            if (r,c) in covered: continue
            cell=ws.cell(r,c)
            x=sum(widths[:c-1])+12
            y=sum(heights[:r-1])-y0+34
            w,h=cell_geometry(ws,r,c,widths,heights,origins)
            fill=color(cell.fill.fgColor) if cell.fill.patternType=='solid' else '#FFFFFF'
            fc=color(cell.font.color,'#242424')
            if ws.title=='11_判定シミュレーター' and c in [4,8] and str(cell.value)=='ALARM':
                fill='#444444'; fc='#FFFFFF'
            draw.rectangle((x,y,x+w,y+h),fill=fill)
            if cell.border.bottom.style:
                draw.line((x,y+h-1,x+w,y+h-1),fill=color(cell.border.bottom.color,'#D9D9D9'))
            if cell.border.left.style: draw.line((x,y,x,y+h),fill='#BBBBBB')
            if cell.border.right.style: draw.line((x+w-1,y,x+w-1,y+h),fill='#BBBBBB')
            if cell.value is None: continue
            f,lines,lh=line_metrics(cell,str(cell.value),w)
            if cell.alignment.vertical in ['center','distributed','justify']: ty=y+(h-lh*len(lines))/2
            else: ty=y+3
            for text in lines:
                if cell.alignment.horizontal=='center': tx=x+(w-f.getlength(text))/2
                elif cell.alignment.horizontal=='right': tx=x+w-f.getlength(text)-5
                else: tx=x+6
                draw.text((tx,ty),text,font=f,fill=fc)
                if cell.font.underline:
                    draw.line((tx,ty+lh-2,tx+f.getlength(text),ty+lh-2),fill=fc)
                ty+=lh
    if ws.title=='01_構成図':
        diagram=Image.open(VALIDATION/'architecture.png').convert('RGB')
        diagram=diagram.resize((round(diagram.width*.82),round(diagram.height*.82)),Image.Resampling.LANCZOS)
        canvas.paste(diagram,(12,34+sum(heights[:5])-y0))
    canvas.save(VALIDATION/name)

def make_previews():
    book=openpyxl.load_workbook(XLSX,data_only=True)
    targets=[('00_読む順番','preview_cover.png',1,24),('01_構成図','preview_architecture.png',1,36),
             ('09_Alarm設定値','preview_alarm.png',1,15),('11_判定シミュレーター','preview_simulator.png',1,28),
             ('16_JP1への連携','preview_jp1.png',1,16),('30_設定例とコマンド','preview_configuration.png',1,40)]
    for sheet,name,start,end in targets: render(book[sheet],name,start,end)
    issues=inspect_layout(book)
    book.close()
    return issues

if __name__=='__main__':
    import json,sys
    sys.stdout.reconfigure(encoding='utf-8')
    print(json.dumps(make_previews(),ensure_ascii=False,indent=2))
